"""
One-off extraction + fuzzy-matching pass for the Maintenance department's
historical data migration ("pls help copy those old data to new database").

Reads the three source Excel files and produces a single JSON file
(migration_data.json) with three arrays ready for a Node upload script:
  - equipment            -> maintenance_equipment
  - inspections          -> maintenance_inspections (keyed by equip_code)
  - pm_schedule          -> maintenance_pm_schedule (keyed by equip_code)

Matching approach (per PRD §7 — "combine it if possible... flag rows that
don't confidently match for a manual pass" rather than silently guess):
  - INVENTORY.xlsx's 680 rows are the primary equipment registry (1 row = 1
    physical item, EquipID is the natural key). No matching needed here.
  - ASSET OWNERSHIP.xlsx's 43 asset-type x location rows are applied onto
    matching INVENTORY equipment via a token-overlap score (asset-type
    keywords like tonnage/brand + a location keyword). Only applied above a
    confidence threshold; everything else is left alone and reported.
  - 2026 PM SCHEDULE.xlsx's 38 named machines are matched the same way
    against INVENTORY; unmatched schedule rows become new minimal equipment
    stubs (named machine, no other fields) so the schedule isn't lost.

Prints a match/no-match report to stdout for manual review before anything
is written to Supabase.
"""
import json
import re
import openpyxl
from datetime import datetime, date

FOLDER = 'C:/Users/hp/Downloads'


def norm(s):
    if s is None:
        return ''
    return re.sub(r'\s+', ' ', str(s)).strip()


def tokens(s):
    """Distinctive tokens: tonnages (40ton), brand words, line/yard numbers."""
    s = norm(s).lower()
    s = s.replace('multicrane', 'multi').replace('straddle crane', 'straddle')
    # Order matters: more specific numeric patterns (tonnage/HP/No./line/
    # yard/N-code) must be tried before the bare \d+ fallback, or e.g.
    # "Shovel Lonking 1" and "Shovel Lonking 2" would tokenize identically
    # (the trailing digit is otherwise dropped entirely, since a lone digit
    # matches none of the specific alternatives) and collide onto the same
    # equipment row.
    toks = set(re.findall(r'\d+\.?\d*ton|\d+hp|no\.?\s?\d+|line\s?\d+|yard\s?[a-z]|n\d+|[a-z]+|\d+', s))
    # keep only meaningful tokens (drop tiny stopword-ish fragments)
    stop = {'the', 'of', 'a', 'and', 'no', 'na', 'area', 'plo', 'plo68', 'plo79'}
    # Keep single-digit tokens (they're often the ONLY thing distinguishing
    # e.g. "Shovel Lonking 1" from "...2", or "Straddle Crane SC1" from
    # "SC2") — only drop single-character alphabetic noise.
    return {t for t in toks if t not in stop and (len(t) > 1 or t.isdigit())}


def score(a_tokens, b_tokens):
    if not a_tokens or not b_tokens:
        return 0
    return len(a_tokens & b_tokens)


def non_digit(tok_set):
    return {t for t in tok_set if not t.isdigit()}


def has_real_overlap(a_tokens, b_tokens):
    """A bare digit (e.g. the trailing '1' in both 'RTG 1' and 'Water Jet
    200bar 1') is not a meaningful asset-type signal on its own — it only
    disambiguates within a type, and two unrelated items sharing a location
    plus a coincidental trailing digit is not a real match. Require at
    least one non-numeric token in common before a digit is allowed to help
    tie-break."""
    return len(non_digit(a_tokens) & non_digit(b_tokens)) >= 1


def to_date_str(v):
    """Handles both python datetime objects and 'M/D/YYYY' strings from the sheet."""
    if v is None or v == '':
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})$', s)
    if m:
        mm, dd, yyyy = m.groups()
        return f'{yyyy}-{mm.zfill(2)}-{dd.zfill(2)}'
    return None


COND_MAP = {'good': 'good', 'fair': 'fair', 'poor': 'poor', 'spoil': 'spoil'}

# ---------------- INVENTORY.xlsx : primary equipment registry ----------------
wb = openpyxl.load_workbook(f'{FOLDER}/INVENTORY.xlsx', data_only=True)
ws = wb.active
headers = [c.value for c in ws[1]]
inv_rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in r):
        continue
    inv_rows.append(dict(zip(headers, r)))

equipment = []
inspections = []
for r in inv_rows:
    cond = COND_MAP.get(norm(r.get('Equipment Condition')).lower())
    equip_code = norm(r.get('EquipID'))
    eq = {
        'equip_code': equip_code,
        'name': norm(r.get('Item')) or equip_code,
        'category': norm(r.get('Category')) or None,
        'brand': norm(r.get('Brand')) or None,
        'purpose': norm(r.get('Purpose')) or None,
        'location': norm(r.get('Location')) or None,
        'condition': cond,
        'manager': None, 'supervisor': None, 'pic_day': None, 'pic_night': None,
        # Category is only a useful match signal when it's a distinctive
        # type (Crane/Compressor/Transport/...) — "Tool" and "Equipment"
        # cover hundreds of unrelated items and, included here, matched
        # e.g. "Cutting Torch" (category "Equipment") to the ownership
        # sheet's "Lifting Equipment"/"Other SBG Equipment" rows purely on
        # that generic word. Leave those two out of the match tokens.
        '_match_tokens': tokens(r.get('Item')) | (
            tokens(r.get('Category')) if norm(r.get('Category')).lower() not in ('tool', 'equipment') else set()
        ),
        '_loc_tokens': tokens(r.get('Location')),
    }
    equipment.append(eq)

    insp_date = to_date_str(r.get('Date Inspect'))
    if insp_date:
        inspections.append({
            'equip_code': equip_code,
            'inspected_by': norm(r.get('Issue By')) or None,
            'inspection_date': insp_date,
            'condition': cond,
            'remarks': norm(r.get('Remarks')) or None,
            'preventive_action_recommendation': norm(r.get('Preventive Action Recommendation')) or None,
        })

print(f'INVENTORY: {len(equipment)} equipment rows, {len(inspections)} inspection rows with a date')

# ---------------- ASSET OWNERSHIP.xlsx : forward-fill + apply ----------------
wb2 = openpyxl.load_workbook(f'{FOLDER}/ASSET OWNERSHIP.xlsx', data_only=True)
ws2 = wb2.active
ownership_rows = []
last_asset = None
for row in ws2.iter_rows(min_row=3, max_row=45, values_only=True):
    asset, loc, mgr, sup, pic_day, pic_night = row[:6]
    if asset:
        last_asset = norm(asset)
    if not loc:
        continue
    ownership_rows.append({
        'asset': last_asset, 'location': norm(loc),
        'manager': norm(mgr) or None, 'supervisor': norm(sup) or None,
        'pic_day': norm(pic_day) or None, 'pic_night': norm(pic_night) or None,
    })

print(f'ASSET OWNERSHIP: {len(ownership_rows)} rows (forward-filled)')

# ASSET OWNERSHIP and INVENTORY use two different location vocabularies for
# the same physical places (e.g. ownership's "N111" vs inventory's
# "PLO 68 N111"). A bare token-overlap check on location text alone misses
# almost all of these. Hand-verified crosswalk from the two location lists
# actually observed in the sheets, only for pairs that are an unambiguous
# match to one specific inventory location:
LOCATION_ALIASES = {
    'n111': {'n111'},
    'fabrication plo79': {'fabrication'},
    'green dot area': {'green', 'dot'},
    'batching plant 1/2': {'batching', 'plant'},
    'batching plant': {'batching', 'plant'},
    'casting plo68': {'production'},  # casting happens in PLO68's "Production" location
    'line 3': {'line 3'},
    'line 4': {'line 4'},
    # deliberately NOT mapped (genuinely ambiguous — no single inventory
    # location corresponds): 'Line 1/2' (covers 2 lines), 'Fabrication PLO68',
    # 'Maintenance', 'Transport Workshop', 'PLO68' (bare), 'External'.
}


def loc_alias_tokens(location_text):
    return LOCATION_ALIASES.get(norm(location_text).lower(), set())


# A candidate is only accepted when BOTH the asset-type tokens AND the
# location tokens overlap with the equipment row (>=1 each) — matching on
# location alone (e.g. two unrelated items that both sit at "PLO 79
# Fabrication") produced false positives in review (a cutting torch getting
# tagged with the "Lifting/Electrical Fitting" owner purely because they
# share a location). Requiring both keeps only matches that are actually
# about the same kind of asset.
applied = 0
ownership_unmatched = []
for o in ownership_rows:
    a_tok = tokens(o['asset'])
    l_tok = tokens(o['location']) | loc_alias_tokens(o['location'])
    best, best_score = None, 0
    for eq in equipment:
        a_s = score(a_tok, eq['_match_tokens'])
        l_s = score(l_tok, eq['_loc_tokens'])
        if a_s < 1 or l_s < 1 or not has_real_overlap(a_tok, eq['_match_tokens']):
            continue
        s = a_s + l_s
        if s > best_score:
            best, best_score = eq, s
    if best is not None:
        best['manager'] = o['manager']
        best['supervisor'] = o['supervisor']
        best['pic_day'] = o['pic_day']
        best['pic_night'] = o['pic_night']
        applied += 1
    else:
        ownership_unmatched.append((o['asset'], o['location'], best_score, best['name'] if best else None))

print(f'ASSET OWNERSHIP applied to {applied}/{len(ownership_rows)} equipment rows')
print(f'Unmatched ownership rows ({len(ownership_unmatched)}) — left blank on equipment, needs manual review:')
for asset, loc, s, guess in ownership_unmatched:
    print(f'  score={s}  asset={asset!r:30s} loc={loc!r:20s} best-guess-was={guess!r}')

# ---------------- 2026 PM SCHEDULE.xlsx : match + build calendar ----------------
wb3 = openpyxl.load_workbook(f'{FOLDER}/2026 PM SCHEDULE.xlsx', data_only=True)
ws3 = wb3['2026 Maintenance Schedule']

YELLOW = 'FFFFFF00'
pm_names_rows = []  # (row_idx, name)
for row in ws3.iter_rows(min_row=3, max_row=45, min_col=1, max_col=1):
    v = row[0].value
    if v:
        pm_names_rows.append((row[0].row, norm(v)))

print(f'\nPM SCHEDULE: {len(pm_names_rows)} named machines')

new_stub_equipment = []
pm_schedule = []
pm_unmatched = []
used_equip_codes = set()  # each machine gets at most one PM-schedule name; a
                           # second name scoring onto the same equipment (e.g.
                           # "MULTICRANE 5ton (Rebar)" and "...(A)" both
                           # matching E00092) is a real distinct machine that
                           # just wasn't distinguishable by name+location
                           # tokens — falls back to a stub rather than
                           # colliding two machines' schedules together.
for row_idx, name in pm_names_rows:
    n_tok = tokens(name)
    best, best_score = None, 0
    for eq in equipment:
        if eq['equip_code'] in used_equip_codes:
            continue
        name_s = score(n_tok, eq['_match_tokens'])
        loc_s = score(n_tok, eq['_loc_tokens'])
        if name_s < 1 or not has_real_overlap(n_tok, eq['_match_tokens']):
            continue  # the PM schedule's own name has to match the equipment name itself
        s = name_s + loc_s
        if s > best_score:
            best, best_score = eq, s
    if best is not None and best_score >= 2:
        equip_code = best['equip_code']
        used_equip_codes.add(equip_code)
    else:
        # No confident match -> create a minimal stub so the schedule isn't lost.
        stub_code = f'PM-{re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-").upper()}'
        if not any(s['equip_code'] == stub_code for s in new_stub_equipment):
            new_stub_equipment.append({
                'equip_code': stub_code, 'name': name, 'category': None, 'brand': None,
                'purpose': None, 'location': None, 'condition': None,
                'manager': None, 'supervisor': None, 'pic_day': None, 'pic_night': None,
            })
        equip_code = stub_code
        pm_unmatched.append((name, best_score, best['name'] if best else None))

    # WK1..WK52 are columns 2..53 (col B onward), one column per week.
    for col in range(2, 54):
        cell = ws3.cell(row=row_idx, column=col)
        fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None
        if fill == YELLOW:
            pm_schedule.append({'equip_code': equip_code, 'year': 2026, 'week_number': col - 1, 'planned': True})

print(f'PM SCHEDULE matched {len(pm_names_rows) - len(pm_unmatched)}/{len(pm_names_rows)} machines to existing equipment')
print(f'PM SCHEDULE created {len(new_stub_equipment)} new stub equipment for unmatched machines:')
for name, s, guess in pm_unmatched:
    print(f'  score={s}  name={name!r:35s} best-guess-was={guess!r}')
print(f'Total planned PM week-cells: {len(pm_schedule)}')

# ---------------- write output ----------------
for eq in equipment:
    eq.pop('_match_tokens', None)
    eq.pop('_loc_tokens', None)

out = {
    'equipment': equipment + new_stub_equipment,
    'inspections': inspections,
    'pm_schedule': pm_schedule,
}
with open('C:/Users/hp/Documents/RebarPro/scripts/maintenance-migration/migration_data.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, indent=2, ensure_ascii=False)

print(f'\nWrote migration_data.json: {len(out["equipment"])} equipment, {len(out["inspections"])} inspections, {len(out["pm_schedule"])} pm_schedule rows')
